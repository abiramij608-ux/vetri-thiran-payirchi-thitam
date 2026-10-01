from io import BytesIO
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Inches, Pt
from fpdf import FPDF
from html import escape


def format_txt(content: str) -> bytes:
    return content.strip().encode("utf-8")


def format_docx(
    content: str,
    document_type: str,
    logo_bytes: Optional[bytes] = None,
) -> bytes:
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if logo_bytes:
        logo_stream = BytesIO(logo_bytes)
        paragraph = document.add_paragraph()
        paragraph.alignment = 1
        run = paragraph.add_run()
        run.add_picture(logo_stream, width=Inches(0.8))

    title = document.add_paragraph()
    title.alignment = 1

    run = title.add_run(document_type)
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph()

    for line in content.splitlines():
        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        if (
            line[:2].isdigit()
            or line[:3].isdigit()
            or line.startswith("SECTION")
            or line.startswith("ARTICLE")
        ):
            run = paragraph.add_run(line)
            run.bold = True
        else:
            paragraph.add_run(line)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class LegalPDF(FPDF):
    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 10, "LegalEase", align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")


def format_pdf(
    content: str,
    document_type: str,
    logo_path: Optional[str] = None,
) -> bytes:
    pdf = LegalPDF()
    pdf.add_page()

    if logo_path and Path(logo_path).exists():
        try:
            pdf.image(logo_path, x=95, y=20, w=20)
            pdf.ln(18)
        except Exception:
            pass

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, document_type, align="C")
    pdf.ln(5)

    pdf.set_font("Helvetica", size=11)

    for line in content.splitlines():
        line = line.strip()

        if not line:
            pdf.ln(4)
            continue

        pdf.multi_cell(0, 7, line)

    return bytes(pdf.output())


def format_html_preview(content: str) -> str:
    html_lines = []

    for line in content.splitlines():
        line = line.strip()

        if not line:
            html_lines.append("<br>")
            continue

        safe_line = escape(line)

        if (
            line[:2].isdigit()
            or line[:3].isdigit()
            or line.startswith("SECTION")
            or line.startswith("ARTICLE")
        ):
            html_lines.append(
                f'<div class="section-title">{safe_line}</div>'
            )
        elif line.startswith(("-", "•", "*")):
            html_lines.append(
                f'<div class="bullet">{safe_line}</div>'
            )
        else:
            html_lines.append(
                f"<p>{safe_line}</p>"
            )

    return "\n".join(html_lines)